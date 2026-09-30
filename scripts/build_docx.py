"""Render the assembled manuscript to a Word .docx.

pandoc is not available here, so this walks the markdown directly. It handles
the subset the manuscript actually uses: ATX headings, paragraphs, bullet and
numbered lists, pipe tables, blockquotes, and inline emphasis/code/math.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# the small LaTeX subset the manuscript uses, mapped to plain unicode
TEX = {
    r"\omega": "ω", r"\tau": "τ", r"\lambda": "λ", r"\beta": "β",
    r"\eta": "η", r"\ell": "ℓ", r"\times": "×", r"\leq": "≤",
    r"\geq": "≥", r"\le": "≤", r"\ge": "≥", r"\neq": "≠",
    r"\to": "→", r"\pm": "±", r"\in": "∈", r"\sum": "∑",
    r"\partial": "∂", r"\top": "⊤", r"\propto": "∝",
    r"\approx": "≈", r"\cdot": "·", r"\qquad": "   ", r"\,": " ",
    r"\;": " ", r"\!": "", r"\%": "%", r"\mathrm": "", r"\text": "",
    r"\mathbf": "", r"\frac": "", r"\min": "min", r"\max": "max", r"\exp": "exp",
}
SUB = str.maketrans("0123456789+-=()aeioxjn", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑᵢₒₓⱼₙ")
SUP = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")


def demath(s: str) -> str:
    def one(m):
        t = m.group(1)
        for k, v in sorted(TEX.items(), key=lambda kv: -len(kv[0])):
            t = t.replace(k, v)
        t = re.sub(r"_\{([^{}]*)\}", lambda x: x.group(1).translate(SUB), t)
        t = re.sub(r"\^\{([^{}]*)\}", lambda x: x.group(1).translate(SUP), t)
        t = re.sub(r"_([A-Za-z0-9])", lambda x: x.group(1).translate(SUB), t)
        t = re.sub(r"\^([A-Za-z0-9])", lambda x: x.group(1).translate(SUP), t)
        return t.replace("{", "").replace("}", "").strip()
    # (?<!\\) so an escaped \$ (a currency figure) is not read as a math delimiter;
    # without it the match runs from the currency sign to the next real $ and eats
    # the prose between them.
    s = re.sub(r"(?<!\\)\$\$(.+?)\$\$", one, s, flags=re.S)
    s = re.sub(r"(?<!\\)\$([^$\n]+?)\$", one, s)
    return s.replace("\\$", "$")


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def runs(par, text, base_bold=False):
    """Inline markdown: **bold**, *italic*, `code`."""
    for part in re.split(r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*"):
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            r = par.add_run(part)
        if base_bold:
            r.bold = True


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", default="paper/phisharmor_manuscript.md")
    ap.add_argument("--out", default="paper/PhishArmor_paper.docx")
    a = ap.parse_args()

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.10
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(0.9)

    lines = Path(a.source).read_text().splitlines()
    i, n = 0, len(lines)
    while i < n:
        ln = lines[i]

        if not ln.strip() or ln.strip() in ("---", "***"):
            i += 1; continue

        # pipe table
        if ln.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i+1]):
            hdr = [c.strip() for c in ln.strip().strip("|").split("|")]
            i += 2
            body = []
            while i < n and lines[i].lstrip().startswith("|"):
                body.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
            t = doc.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for k, h in enumerate(hdr):
                c = t.rows[0].cells[k]; c.text = ""
                runs(c.paragraphs[0], demath(h), base_bold=True)
                c.paragraphs[0].runs and setattr(c.paragraphs[0].runs[0].font, "size", Pt(9))
                shade(c, "EFF1F4")
            for row in body:
                cells = t.add_row().cells
                for k, v in enumerate(row[:len(hdr)]):
                    cells[k].text = ""
                    runs(cells[k].paragraphs[0], demath(v))
                    for r in cells[k].paragraphs[0].runs: r.font.size = Pt(9)
            doc.add_paragraph()
            continue

        # heading
        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            lvl, txt = len(m.group(1)), demath(m.group(2)).strip()
            if lvl == 1 and txt.lower().startswith("phisharmor"):
                p = doc.add_heading(txt, 0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_heading(txt, min(lvl, 4))
            i += 1; continue

        # blockquote
        if ln.lstrip().startswith(">"):
            buf = []
            while i < n and lines[i].lstrip().startswith(">"):
                buf.append(lines[i].lstrip()[1:].strip()); i += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            runs(p, demath(" ".join(buf)))
            for r in p.runs: r.italic = True; r.font.color.rgb = RGBColor(0x55, 0x5B, 0x66)
            continue

        # html block (byline) -> centred plain lines
        if ln.lstrip().startswith("<div"):
            while i < n and "</div>" not in lines[i]:
                txt = re.sub(r"<[^>]+>", "", lines[i]).strip()
                if txt:
                    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    runs(p, txt)
                i += 1
            i += 1; continue

        # lists
        m = re.match(r"^(\s*)([-*])\s+(.*)$", ln)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            if len(m.group(1)) >= 2: p.paragraph_format.left_indent = Inches(0.6)
            runs(p, demath(m.group(3))); i += 1; continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", ln)
        if m:
            p = doc.add_paragraph(style="List Number")
            runs(p, demath(m.group(3))); i += 1; continue

        # paragraph (join soft-wrapped lines)
        buf = [ln.strip()]; i += 1
        while i < n and lines[i].strip() and not re.match(r"^(#{1,4}\s|\s*[-*]\s|\s*\d+\.\s|>|\||---)", lines[i]):
            buf.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        runs(p, demath(" ".join(buf)))

    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(a.out)
    print(f"wrote {a.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
